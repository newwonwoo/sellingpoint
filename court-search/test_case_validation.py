"""법원·OCR·artifact 통신 없이 실제 배치와 파일 저장 경로를 검증한다."""
import asyncio
import contextlib
import io
import json
import os
from pathlib import Path
import tempfile
import unittest
from unittest.mock import AsyncMock, MagicMock, patch

from test_search_scheduler import SEARCH
import openpyxl
import output_schema


class CaseValidationTests(unittest.TestCase):
    def test_exact_ascii_year_and_complete_case_number(self):
        for value in ('202타인4330', '20266타인4330', '', None, 2026,
                      True, [], {}, '２０２６타인4330', '٢٠٢٦타인4330',
                      '2026타인4330junk', '2026타인４３３０'):
            with self.subTest(value=value), self.assertRaisesRegex(ValueError, 'ASCII 숫자 4자리'):
                SEARCH.parse_case_no(value)
        for value in ('2026타인4330', ' 2026타인4330 ', '1999타경1'):
            self.assertEqual(SEARCH.parse_case_no(value)[0], value.strip()[:4])

    def test_raw_json_and_excel_values_survive_reading(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / 'input.json'
            values = ['202타인4330', '', None, 2026, {}, ' 2026타인4330 ']
            path.write_text(json.dumps([{'court': '서울남부지방법원', 'case_no': v} for v in values]))
            self.assertEqual([c['case_no'] for c in SEARCH.read_input(path)], values)
            path = Path(tmp) / 'input.xlsx'
            wb = openpyxl.Workbook()
            wb.active.append(['법원', '사건번호'])
            for value in values[:4]:
                wb.active.append(['서울남부지방법원', value])
            wb.save(path)
            wb.close()
            self.assertEqual([c['case_no'] for c in SEARCH.read_input(path)], [values[0], None, None, 2026])

    def test_failed_save_preserves_previous_file(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / 'result.xlsx'
            path.write_bytes(b'previous result')
            with patch.object(output_schema, '_write_output', side_effect=OSError('disk full')):
                with self.assertRaises(OSError):
                    output_schema.write_output([], path)
            self.assertEqual(path.read_bytes(), b'previous result')
            self.assertEqual(list(Path(tmp).iterdir()), [path])

    def test_failed_atomic_replace_preserves_previous_file(self):
        with tempfile.TemporaryDirectory() as tmp:
            for suffix in ('xlsx', 'json'):
                path = Path(tmp) / f'result.{suffix}'
                path.write_bytes(b'previous result')
                with patch.object(output_schema.os, 'replace', side_effect=OSError('replace denied')):
                    with self.assertRaises(OSError):
                        output_schema.write_output([{'사건명': '정상'}], path)
                self.assertEqual(path.read_bytes(), b'previous result')
                path.unlink()
                self.assertEqual(list(Path(tmp).iterdir()), [])


class BatchIsolationTests(unittest.IsolatedAsyncioTestCase):
    async def test_batch_continues_and_saves_partials_and_final(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source, output = root / 'input.json', root / 'output.xlsx'
            values = ['2026타인1', '202타인4330', '2026타인2', '2026타인3',
                      None, {}, '２０２６타인4']
            source.write_text(json.dumps([{'court': '서울남부지방법원', 'case_no': v} for v in values]))
            original = source.read_bytes()
            page = MagicMock()
            context = MagicMock(new_page=AsyncMock(return_value=page))
            browser = MagicMock(new_context=AsyncMock(return_value=context), close=AsyncMock())
            playwright = MagicMock()
            playwright.chromium.launch = AsyncMock(return_value=browser)
            manager = MagicMock(__aenter__=AsyncMock(return_value=playwright), __aexit__=AsyncMock())
            calls = []

            async def search_one(page, court, case_no, *args):
                calls.append(case_no)
                if case_no == '2026타인2':
                    raise RuntimeError('mock row failure')
                return {'법원': court, '사건번호': case_no, '사건명': '정상'}

            async def offline_upload(publisher):
                while True:
                    item = await publisher.queue.get()
                    publisher.queue.task_done()
                    if item is None:
                        return

            env = {'COURT_WORKERS': '1', 'COURT_MAX_WORKERS': '1',
                   'COURT_PARTIAL_DIR': str(root / 'partial'),
                   'COURT_ARTIFACT_QUEUE_DIR': str(root / 'queue'), 'COURT_PARTIAL_SIZE': '2'}
            log = io.StringIO()
            with patch.dict(os.environ, env, clear=True), \
                 patch.object(SEARCH.sys, 'argv', ['search.py', str(source), str(output)]), \
                 patch.object(SEARCH, 'async_playwright', return_value=manager), \
                 patch.object(SEARCH, 'search_one', side_effect=search_one), \
                 patch.object(SEARCH, 'load_success_dataset', return_value={'loaded': 0, 'skipped': 0}), \
                 patch.object(SEARCH, 'FAST_CAPTCHA_PATH', False), \
                 patch.object(SEARCH.PartialResultPublisher, '_upload_worker', offline_upload), \
                 patch('socket.socket.connect', side_effect=AssertionError('network forbidden')), \
                 contextlib.redirect_stdout(log):
                await SEARCH.main()
            self.assertEqual(calls, ['2026타인1', '2026타인2', '2026타인3'])
            self.assertEqual(source.read_bytes(), original)
            wb = openpyxl.load_workbook(output)
            rows = list(wb.active.values)[1:]
            wb.close()
            self.assertEqual(len(rows), len(values))
            self.assertEqual(rows[1][1], '202타인4330')
            self.assertIn('입력오류', rows[1][2])
            self.assertIn('RuntimeError: mock row failure', rows[2][2])
            self.assertEqual(rows[3][2], '정상')
            self.assertIn('ASCII 숫자 4자리', log.getvalue())
            partial_rows = []
            for path in sorted((root / 'partial').glob('*.xlsx')):
                wb = openpyxl.load_workbook(path)
                partial_rows.extend(list(wb.active.values)[1:])
                wb.close()
            self.assertEqual(partial_rows, rows)

    async def test_all_invalid_never_starts_browser_and_saves_json(self):
        with tempfile.TemporaryDirectory() as tmp:
            source, output = Path(tmp) / 'input.json', Path(tmp) / 'output.json'
            source.write_text(json.dumps({'court': '서울남부지방법원', 'case_no': None}))
            with patch.dict(os.environ, {}, clear=True), \
                 patch.object(SEARCH.sys, 'argv', ['search.py', str(source), str(output)]), \
                 patch.object(SEARCH, 'async_playwright', side_effect=AssertionError('browser forbidden')):
                await SEARCH.main()
            data = json.loads(output.read_text())['results']
            self.assertIsNone(data[0]['사건번호'])
            self.assertIn('입력오류', data[0]['사건명'])

    async def test_cancellation_is_not_swallowed(self):
        case = {'row': 2, 'court': '서울남부지방법원', 'case_no': '2026타인1', 'party_name': '주택'}
        with patch.object(SEARCH, 'search_one', side_effect=asyncio.CancelledError):
            with self.assertRaises(asyncio.CancelledError):
                await SEARCH.search_case_safely(None, case)


if __name__ == '__main__':
    unittest.main()
