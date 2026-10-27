"""workflow의 Python 검증 코드만 임시 XLSX에서 실행한다(CI/통신 없음)."""
from contextlib import redirect_stdout
import io
import os
from pathlib import Path
import tempfile
import textwrap
import unittest
from unittest.mock import patch

import openpyxl
import chunk_workbook
from test_search_scheduler import SEARCH


class WorkflowInputTests(unittest.TestCase):
    def test_count_blocked_output_and_retry_preserve_invalid_rows(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source, output = root / 'input.xlsx', root / 'output.xlsx'
            wb = openpyxl.Workbook()
            wb.active.append(['법원', '사건번호'])
            values = ['2026타인1', None, 0, '202타인4330', '2026타인2']
            for value in values:
                wb.active.append(['서울남부지방법원', value])
            wb.save(source)
            wb.close()
            header, rows = chunk_workbook.input_rows(source)
            self.assertEqual(len(rows), len(SEARCH.read_input(source)))
            self.assertEqual(len(rows), 5)
            chunk_workbook.write_blocked_output(source, output, '법원 사이트 접속 불가')
            wb = openpyxl.load_workbook(output)
            results = list(wb.active.values)[1:]
            wb.close()
            self.assertEqual([r[1] for r in results], values)
            self.assertTrue(all(r[2].startswith('입력오류') for r in results[1:4]))
            retry, count = root / 'retry.xlsx', root / 'count.txt'
            chunk_workbook.write_retry_input(source, output, retry, count)
            self.assertEqual(count.read_text(), '2')
            self.assertEqual([c['case_no'] for c in SEARCH.read_input(retry)], ['2026타인1', '2026타인2'])
            # 실제 workflow inspect 코드의 건수 검증/실패 집계/fallback 판단을 실행.
            workflow = Path(__file__).parents[1] / '.github/workflows/court_search.yml'
            text = workflow.read_text()
            for name in ('결과 파일 생성 확인', '재조회 결과 파일 생성 확인'):
                section = text.split(f'- name: {name}', 1)[1]
                code = textwrap.dedent(section.split("<<'PY'\n", 1)[1].split('\n          PY', 1)[0])
                code = code.replace("Path('court-search/output.xlsx')", f'Path({str(output)!r})')
                github_output = root / 'github-output.txt'
                with patch.dict(os.environ, {'EXPECTED_TOTAL': '5', 'GITHUB_OUTPUT': str(github_output)}), redirect_stdout(io.StringIO()):
                    namespace = {}
                    exec(compile(code, str(workflow), 'exec'), namespace)
                self.assertEqual(namespace['actual'], 5)
                self.assertEqual(len(namespace['failed']), 5)
                if name == '결과 파일 생성 확인':
                    self.assertTrue(namespace['global_connection_failure'])
                    self.assertIn('global_connection_failure=true', github_output.read_text())

    def test_count_and_blocked_output_use_header_names(self):
        with tempfile.TemporaryDirectory() as tmp:
            source, output = Path(tmp) / 'input.xlsx', Path(tmp) / 'output.xlsx'
            wb = openpyxl.Workbook()
            wb.active.append(['사건번호', '메모', '법원'])
            wb.active.append([None, '빈 사건번호 유지', '서울남부지방법원'])
            wb.active.append([None, None, None])
            wb.active.append(['2026타인1', None, '서울남부지방법원'])
            wb.save(source)
            wb.close()
            self.assertEqual(len(chunk_workbook.input_rows(source)[1]), 2)
            self.assertEqual(len(SEARCH.read_input(source)), 2)
            chunk_workbook.write_blocked_output(source, output, '법원 사이트 접속 불가')
            wb = openpyxl.load_workbook(output)
            rows = list(wb.active.values)[1:]
            wb.close()
            self.assertIsNone(rows[0][1])
            self.assertTrue(rows[0][2].startswith('입력오류'))
            self.assertEqual(rows[1][1], '2026타인1')


if __name__ == '__main__':
    unittest.main()
