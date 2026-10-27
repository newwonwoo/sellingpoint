"""사건조회 JSON·엑셀 결과의 공통 스키마와 출력 함수."""

from __future__ import annotations

import json
import os
import tempfile
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill


OUTPUT_HEADERS = [
    '법원', '사건번호', '사건명', '재판부',
    '접수일', '종국결과', '결정문송달일', '확정일',
    '진행_1일자', '진행_1내용', '진행_1결과', '진행_1공시문',
    '진행_2일자', '진행_2내용', '진행_2결과', '진행_2공시문',
    '진행_3일자', '진행_3내용', '진행_3결과', '진행_3공시문',
    '관련사건_법원', '관련사건_번호',
    '신청인', '피신청인', '조회일시',
]


def write_output(results: list[dict], path: str):
    """완성된 파일만 교체하여 저장 중 오류가 기존 결과를 손상시키지 않는다."""
    target = Path(path)
    fd, temporary = tempfile.mkstemp(dir=target.parent, suffix=target.suffix)
    os.close(fd)
    try:
        _write_output(results, temporary)
        os.replace(temporary, target)
    finally:
        Path(temporary).unlink(missing_ok=True)
    print(f"\n✅ 결과 저장: {path} ({len(results)}건)")


def _write_output(results: list[dict], path: str):
    if Path(path).suffix.lower() == '.json':
        normalized = [
            {header: data.get(header, '') for header in OUTPUT_HEADERS}
            for data in results
        ]
        Path(path).write_text(
            json.dumps({'results': normalized}, ensure_ascii=False, indent=2),
            encoding='utf-8',
        )
        return

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = '조회결과'

    header_fill = PatternFill(
        start_color='1F4E79', end_color='1F4E79', fill_type='solid'
    )
    header_font = Font(bold=True, color='FFFFFF', size=10)

    for column, header in enumerate(OUTPUT_HEADERS, 1):
        cell = sheet.cell(row=1, column=column, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center')

    for row_number, data in enumerate(results, 2):
        for column, header in enumerate(OUTPUT_HEADERS, 1):
            value = data.get(header, '')
            if isinstance(value, (dict, list)):
                value = json.dumps(value, ensure_ascii=False)
            cell = sheet.cell(row=row_number, column=column, value=value)
            if isinstance(value, str):
                cell.data_type = 's'  # 원 입력을 Excel 수식으로 실행하지 않는다.

    column_widths = {
        1: 20, 2: 16, 3: 24, 4: 20,
        5: 12, 6: 20, 7: 14, 8: 12,
        9: 12, 10: 30, 11: 20, 12: 16,
        13: 12, 14: 30, 15: 20, 16: 16,
        17: 12, 18: 30, 19: 20, 20: 16,
        21: 18, 22: 16,
        23: 24, 24: 24, 25: 18,
    }
    for column, width in column_widths.items():
        sheet.column_dimensions[
            openpyxl.utils.get_column_letter(column)
        ].width = width

    sheet.freeze_panes = 'A2'
    workbook.save(path)
    workbook.close()
