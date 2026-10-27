"""20건 단위 조회 작업용 입력 분할과 결과 병합 도구."""

from __future__ import annotations

import argparse
import json
from copy import copy
from datetime import datetime
from pathlib import Path

import openpyxl
from output_schema import OUTPUT_HEADERS, write_output
from case_number import parse_case_no


CHUNK_SIZE = 20
RETRY_PREFIXES = ('조회실패', '조회보류')


def input_rows(path: Path) -> tuple[list, list[tuple]]:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        raise ValueError("입력 엑셀이 비어 있습니다.")
    workbook.close()
    header = list(rows[0])
    columns = {str(value).strip(): index for index, value in enumerate(header)}
    if '법원' not in columns or '사건번호' not in columns:
        raise ValueError("입력 파일에 '법원', '사건번호' 열이 필요합니다.")
    data = [row for row in rows[1:] if any(value is not None for value in row)]
    if any(not row[columns['법원']] for row in data):
        raise ValueError('모든 입력 행에 법원이 필요합니다.')
    return header, data


def print_matrix(path: Path):
    _, rows = input_rows(path)
    if not rows:
        raise ValueError("조회할 사건이 없습니다.")
    chunks = []
    for start_zero in range(0, len(rows), CHUNK_SIZE):
        start = start_zero + 1
        end = min(start_zero + CHUNK_SIZE, len(rows))
        chunks.append({
            "chunk": len(chunks) + 1,
            "start": start,
            "end": end,
            "count": end - start + 1,
            "tag": f"{start:04d}-{end:04d}",
        })
    workers = 3 if len(chunks) == 1 else 1
    for chunk in chunks:
        chunk["workers"] = workers
    print(f"matrix={json.dumps({'include': chunks}, ensure_ascii=False, separators=(',', ':'))}")
    print(f"total={len(rows)}")
    print(f"chunk_count={len(chunks)}")


def print_count(path: Path):
    _, rows = input_rows(path)
    if not rows:
        raise ValueError("조회할 사건이 없습니다.")
    print(f"total={len(rows)}")


def write_blocked_output(source: Path, target: Path, reason: str):
    """연결 진단 실패 시 실제 조회 없이 모든 입력을 보류 결과로 남긴다."""
    header, rows = input_rows(source)
    columns = {str(value).strip(): index for index, value in enumerate(header)}
    checked_at = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
    results = []
    for row in rows:
        case_no = row[columns['사건번호']]
        status = f'조회보류: {reason}'
        try:
            parse_case_no(case_no)
        except ValueError as error:
            status = f'입력오류: {error}'
            print(f'⚠ 원 사건번호 {case_no!r}: {status}')
        results.append({header: '' for header in OUTPUT_HEADERS} | {
            '법원': str(row[columns['법원']]).strip(),
            '사건번호': case_no,
            '사건명': status,
            '조회일시': checked_at,
        })
    target.parent.mkdir(parents=True, exist_ok=True)
    write_output(results, str(target))
    print(f'⏹ 연결 진단 실패 결과 생성: {len(results)}건')


def write_slice(source: Path, target: Path, start: int, end: int):
    header, rows = input_rows(source)
    selected = rows[start - 1:end]
    if not selected:
        raise ValueError(f"입력 범위가 비어 있습니다: {start}-{end}")

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "입력"
    sheet.append(header)
    for row in selected:
        sheet.append(list(row))
    target.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(target)
    print(f"✅ 입력 묶음 생성: {target} ({start}-{end}, {len(selected)}건)")


def write_retry_input(source: Path, result: Path, target: Path, count_file: Path):
    """실패·보류된 사건만 새 Actions 러너에서 재조회할 입력으로 만든다."""
    header, rows = input_rows(source)
    result_book = openpyxl.load_workbook(result, read_only=True, data_only=True)
    result_sheet = result_book.active
    result_headers = [cell.value for cell in next(result_sheet.iter_rows())]
    court_col = result_headers.index('법원')
    case_col = result_headers.index('사건번호')
    name_col = result_headers.index('사건명')
    retry_keys = {
        (str(row[court_col] or '').strip(), str(row[case_col] or '').strip())
        for row in result_sheet.iter_rows(min_row=2, values_only=True)
        if str(row[name_col] or '').startswith(RETRY_PREFIXES)
    }
    selected = [
        row for row in rows
        if (str(row[0] or '').strip(), str(row[1] or '').strip()) in retry_keys
    ]

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = '재조회입력'
    sheet.append(header)
    for row in selected:
        sheet.append(list(row))
    target.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(target)
    count_file.write_text(str(len(selected)), encoding='utf-8')
    print(f"✅ 새 러너 재조회 대상: {len(selected)}건")


def replace_retry_results(primary: Path, retry: Path, target: Path):
    """새 러너의 재조회 결과로 기존 실패 행만 교체하고 순서를 유지한다."""
    workbook = openpyxl.load_workbook(primary)
    sheet = workbook.active
    headers = [cell.value for cell in sheet[1]]
    court_col = headers.index('법원') + 1
    case_col = headers.index('사건번호') + 1

    retry_book = openpyxl.load_workbook(retry, read_only=True, data_only=True)
    retry_sheet = retry_book.active
    retry_headers = [cell.value for cell in next(retry_sheet.iter_rows())]
    if retry_headers != headers:
        raise ValueError('재조회 결과 열 구성이 기존 결과와 다릅니다.')
    replacements = {
        (str(row[court_col - 1] or '').strip(), str(row[case_col - 1] or '').strip()): row
        for row in retry_sheet.iter_rows(min_row=2, values_only=True)
    }

    replaced = 0
    for row_number in range(2, sheet.max_row + 1):
        key = (
            str(sheet.cell(row_number, court_col).value or '').strip(),
            str(sheet.cell(row_number, case_col).value or '').strip(),
        )
        replacement = replacements.get(key)
        if replacement is None:
            continue
        for column, value in enumerate(replacement, 1):
            sheet.cell(row_number, column, value=value)
        replaced += 1
    if replaced != len(replacements):
        raise ValueError(f'재조회 결과 교체 누락: {replaced}/{len(replacements)}')
    target.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(target)
    print(f"✅ 새 러너 재조회 결과 반영: {replaced}건")


def combine_results(source_dir: Path, target: Path):
    files = sorted(source_dir.glob("court-search-*-*.xlsx"))
    if not files:
        raise ValueError("병합할 중간 결과 파일이 없습니다.")

    first = openpyxl.load_workbook(files[0])
    output = first
    output_sheet = output.active
    total = output_sheet.max_row - 1

    for path in files[1:]:
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
        sheet = workbook.active
        for row in sheet.iter_rows(min_row=2, values_only=True):
            output_sheet.append(list(row))
            total += 1

    # 첫 파일의 열 너비·머리글 스타일을 유지하고 추가 행에는 기본 정렬만 적용한다.
    for cell in output_sheet[1]:
        cell.font = copy(cell.font)
        cell.fill = copy(cell.fill)
        cell.alignment = copy(cell.alignment)
    output_sheet.freeze_panes = "A2"
    target.parent.mkdir(parents=True, exist_ok=True)
    output.save(target)
    print(f"✅ 전체 결과 병합: {target} ({total}건, {len(files)}개 파일)")


def main():
    parser = argparse.ArgumentParser()
    subparsers = parser.add_subparsers(dest="command", required=True)

    matrix_parser = subparsers.add_parser("matrix")
    matrix_parser.add_argument("input", type=Path)

    count_parser = subparsers.add_parser("count")
    count_parser.add_argument("input", type=Path)

    blocked_parser = subparsers.add_parser("blocked-output")
    blocked_parser.add_argument("input", type=Path)
    blocked_parser.add_argument("output", type=Path)
    blocked_parser.add_argument("reason")

    slice_parser = subparsers.add_parser("slice")
    slice_parser.add_argument("input", type=Path)
    slice_parser.add_argument("output", type=Path)
    slice_parser.add_argument("start", type=int)
    slice_parser.add_argument("end", type=int)

    combine_parser = subparsers.add_parser("combine")
    combine_parser.add_argument("source_dir", type=Path)
    combine_parser.add_argument("output", type=Path)

    retry_parser = subparsers.add_parser("retry-input")
    retry_parser.add_argument("input", type=Path)
    retry_parser.add_argument("result", type=Path)
    retry_parser.add_argument("output", type=Path)
    retry_parser.add_argument("count_file", type=Path)

    replace_parser = subparsers.add_parser("replace-retry")
    replace_parser.add_argument("primary", type=Path)
    replace_parser.add_argument("retry", type=Path)
    replace_parser.add_argument("output", type=Path)

    args = parser.parse_args()
    if args.command == "matrix":
        print_matrix(args.input)
    elif args.command == "count":
        print_count(args.input)
    elif args.command == "blocked-output":
        write_blocked_output(args.input, args.output, args.reason)
    elif args.command == "slice":
        write_slice(args.input, args.output, args.start, args.end)
    elif args.command == "combine":
        combine_results(args.source_dir, args.output)
    elif args.command == "retry-input":
        write_retry_input(args.input, args.result, args.output, args.count_file)
    else:
        replace_retry_results(args.primary, args.retry, args.output)


if __name__ == "__main__":
    main()
