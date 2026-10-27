import asyncio
import importlib.util
import os
import tempfile
import unittest
from contextlib import ExitStack
from pathlib import Path
from unittest.mock import AsyncMock, MagicMock, patch

import openpyxl


MODULE_PATH = Path(__file__).with_name('search.py')
SPEC = importlib.util.spec_from_file_location('court_search_runtime', MODULE_PATH)
SEARCH = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(SEARCH)


def sample(*, success=True, connection_failure=False, duration=12.0):
    return {
        'success': success,
        'connection_failure': connection_failure,
        'duration': duration,
    }


class AdaptiveWorkerTests(unittest.TestCase):
    def test_skipped_input_does_not_count_as_site_health_sample(self):
        state = {'enabled': True, 'samples': [], 'scale_event': asyncio.Event(),
                 'extra_worker_retired': False}
        for _ in range(20):
            SEARCH.record_adaptive_health(
                state, {'사건명': '조회보류: 사건번호 형식 오류'}, 0.0
            )
        self.assertEqual(state['samples'], [])
        self.assertFalse(state['scale_event'].is_set())

    def test_scale_after_fifteen_fast_stable_cases(self):
        self.assertFalse(SEARCH.adaptive_scale_allowed([sample()] * 14))
        self.assertTrue(SEARCH.adaptive_scale_allowed([sample()] * 15))

    def test_connection_failure_blocks_scale(self):
        samples = [sample()] * 14 + [sample(success=False, connection_failure=True)]
        self.assertFalse(SEARCH.adaptive_scale_allowed(samples))

    def test_two_recent_connection_failures_retire_extra_worker(self):
        samples = [sample()] * 8 + [
            sample(success=False, connection_failure=True),
            sample(success=False, connection_failure=True),
        ]
        self.assertTrue(SEARCH.adaptive_retire_required(samples))

    def test_recording_health_activates_scale_event(self):
        state = {
            'enabled': True,
            'samples': [],
            'scale_event': asyncio.Event(),
            'extra_worker_retired': False,
        }
        for _ in range(15):
            SEARCH.record_adaptive_health(state, {'사건명': '부동산인도명령'}, 12.0)
        self.assertTrue(state['scale_event'].is_set())


class ConnectionCircuitTests(unittest.IsolatedAsyncioTestCase):
    def circuit(self):
        return {
            'failed_sessions': set(),
            'required_failed_sessions': 3,
            'event': asyncio.Event(),
            # 단위 테스트에서는 외부 DNS/HTTP 진단을 실행하지 않는다.
            'diagnosed': True,
        }

    async def test_retries_from_one_session_do_not_stop_batch(self):
        circuit = self.circuit()
        page = object()
        for _ in range(4):
            await SEARCH.mark_session_connection_failure(
                page, circuit, {'worker_id': 1}
            )
        self.assertEqual(circuit['failed_sessions'], {1})
        self.assertFalse(circuit['event'].is_set())

    async def test_three_distinct_failed_sessions_stop_batch(self):
        circuit = self.circuit()
        page = object()
        for worker_id in (1, 2, 3):
            await SEARCH.mark_session_connection_failure(
                page, circuit, {'worker_id': worker_id}
            )
        self.assertEqual(circuit['failed_sessions'], {1, 2, 3})
        self.assertTrue(circuit['event'].is_set())


class StructureCircuitTests(unittest.TestCase):
    def test_same_case_does_not_open_structure_circuit_twice(self):
        circuit = {}
        SEARCH.register_shared_failure(
            circuit, 'failures', 'event', '2026타인3944', 2
        )
        count, opened = SEARCH.register_shared_failure(
            circuit, 'failures', 'event', '2026타인3944', 2
        )

        self.assertEqual(count, 1)
        self.assertFalse(opened)
        self.assertFalse(circuit['event'].is_set())

    def test_two_distinct_cases_open_structure_circuit(self):
        circuit = {}
        SEARCH.register_shared_failure(
            circuit, 'failures', 'event', '2026타인3944', 2
        )
        count, opened = SEARCH.register_shared_failure(
            circuit, 'failures', 'event', '2026타인3273', 2
        )

        self.assertEqual(count, 2)
        self.assertTrue(opened)
        self.assertTrue(circuit['event'].is_set())


class InvalidCaseNumberTests(unittest.IsolatedAsyncioTestCase):
    async def test_invalid_year_skips_without_page_captcha_or_session_access(self):
        class ForbiddenPage:
            def __getattr__(self, name):
                raise AssertionError(f'조회하면 안 되는 입력에서 페이지 접근: {name}')

        for case_no in ('202타인4330', '26타인4330', '20266타인4330', '타인4330'):
            with self.subTest(case_no=case_no):
                session = {'form_ready': True, 'navigations': 5}
                result = await SEARCH.search_one(
                    ForbiddenPage(), '서울남부지방법원', case_no,
                    connection_circuit=object(), session_state=session,
                )
                self.assertEqual(result['사건번호'], case_no)
                self.assertEqual(result['사건명'], '조회보류: 사건번호 형식 오류')
                self.assertEqual(set(result), set(SEARCH.OUTPUT_HEADERS))
                self.assertEqual(session, {'form_ready': True, 'navigations': 5})
                self.assertFalse(result['사건명'].startswith('조회실패'))

    def test_normal_four_digit_year_keeps_existing_parsing(self):
        self.assertEqual(SEARCH.parse_case_no('2026타인4330'), ('2026', '타인', '4330'))
        self.assertEqual(SEARCH.parse_case_no('1999가단123'), ('1999', '가단', '123'))

    async def test_batch_continues_after_bad_row_and_saves_all_rows_once(self):
        """실제 입력/출력·스케줄러를 실행하고 정상 사건의 외부 조회만 모의 처리한다."""
        original_search = SEARCH.search_one
        visited = []
        numbers = ['2026타인4327', '202타인4330', '2026타인4331']

        async def lookup(page, court, case_no, *args):
            visited.append(case_no)
            if case_no == '202타인4330':
                return await original_search(object(), court, case_no)
            return {h: '' for h in SEARCH.OUTPUT_HEADERS} | {
                '법원': court, '사건번호': case_no, '사건명': '[전자]부동산인도명령',
            }

        browser = MagicMock()
        context = MagicMock()
        context.new_page = AsyncMock(return_value=MagicMock())
        browser.new_context = AsyncMock(return_value=context)
        browser.close = AsyncMock()
        playwright = MagicMock()
        playwright.chromium.launch = AsyncMock(return_value=browser)
        manager = MagicMock()
        manager.__aenter__ = AsyncMock(return_value=playwright)
        manager.__aexit__ = AsyncMock(return_value=False)

        with tempfile.TemporaryDirectory() as directory, ExitStack() as stack:
            source = Path(directory) / 'input.xlsx'
            output = Path(directory) / 'output.xlsx'
            workbook = openpyxl.Workbook()
            workbook.active.append(['법원', '사건번호'])
            for number in numbers:
                workbook.active.append(['서울남부지방법원', number])
            workbook.save(source)
            workbook.close()
            stack.enter_context(patch.object(SEARCH.sys, 'argv', ['search.py', str(source), str(output)]))
            stack.enter_context(patch.dict(os.environ, {
                'COURT_WORKERS': '2', 'COURT_MAX_WORKERS': '2',
                'PLAYWRIGHT_CHROMIUM_EXECUTABLE_PATH': '',
                'COURT_PARTIAL_DIR': '', 'COURT_ARTIFACT_QUEUE_DIR': '',
            }))
            stack.enter_context(patch.object(SEARCH, 'async_playwright', return_value=manager))
            stack.enter_context(patch.object(SEARCH, 'search_one', side_effect=lookup))
            stack.enter_context(patch.object(SEARCH, 'load_success_dataset', return_value={'loaded': 0, 'skipped': 0}))
            for flag in ('FAST_CAPTCHA_PATH', 'REUSE_SEARCH_SESSION', 'DIRECT_API_PATH', 'IMMEDIATE_RETRY_WORKER'):
                stack.enter_context(patch.object(SEARCH, flag, False))
            await SEARCH.main()
            saved = openpyxl.load_workbook(output, read_only=True, data_only=True)
            rows = list(saved.active.iter_rows(values_only=True))
            saved.close()
            self.assertEqual(len(rows), 4)
            self.assertEqual([row[1] for row in rows[1:]], numbers)
            self.assertEqual(rows[2][2], '조회보류: 사건번호 형식 오류')
            self.assertEqual(rows[3][2], '[전자]부동산인도명령')
            self.assertCountEqual(visited, numbers)  # 잘못된 행도 재조회하지 않는다.


if __name__ == '__main__':
    unittest.main()
